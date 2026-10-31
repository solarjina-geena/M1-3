"""엑셀·CSV 업로드 -> JournalPayload 목록(미리보기). 한 행 = 작업일보 1장.

열 이름(머리글): model, process, work_date, shift, gross_output_qty, work_minutes, inputs, materials, defects, lots, file, page
  inputs    : "IMEB3_STITCHING:9.2; IMEB3_INJECTION:8.28"   (공정:LOT날짜 — 공정은 생략 가능 → "9.2")
  materials : "WIRE:8.28; LEAD_PIN:7.22; RESIN:26.07.02"
  defects   : "vision1=24; weld=8; Marking=3"
  lots      : "ZH4=2BOX; ZH3=6BOX"   (최종검사)
"""
from __future__ import annotations

import csv
import io
import re
from typing import Iterable

import openpyxl

from .journal_payload import JournalPayload, MATERIAL_CODES

COLUMNS = ["model", "process", "work_date", "shift", "gross_output_qty", "work_minutes", "inputs", "materials", "defects", "lots", "file", "page"]
PROCESS_ALIAS = {"사출": "IMEB3_INJECTION", "INJECTION": "IMEB3_INJECTION", "BB": "IMEB3_INJECTION",
                 "스티칭": "IMEB3_STITCHING", "STITCHING": "IMEB3_STITCHING", "BS": "IMEB3_STITCHING",
                 "권선": "IMEB3_WINDING", "BA": "IMEB3_WINDING", "WINDING": "IMEB3_WINDING",
                 "압입": "IMEB3_CA_PRESS", "CA": "IMEB3_CA_PRESS", "CA_PRESS": "IMEB3_CA_PRESS",
                 "최종검사": "IMEB3_FINAL_INSPECTION", "FINAL": "IMEB3_FINAL_INSPECTION", "FINAL_INSPECTION": "IMEB3_FINAL_INSPECTION"}


def _split(s) -> list[str]:
    return [p.strip() for p in re.split(r"[;\n]", str(s or "")) if p.strip()]


def _pairs(s) -> list[tuple[str | None, str]]:
    out = []
    for p in _split(s):
        if ":" in p:
            k, v = p.split(":", 1)
            out.append((k.strip(), v.strip()))
        elif "=" in p:
            k, v = p.split("=", 1)
            out.append((k.strip(), v.strip()))
        else:
            out.append((None, p))
    return out


def row_to_payload(row: dict, filename: str, method: str) -> JournalPayload:
    g = lambda k: (str(row.get(k)).strip() if row.get(k) is not None else "")
    proc = g("process").upper()
    proc = PROCESS_ALIAS.get(proc, PROCESS_ALIAS.get(g("process"), proc))
    inputs = [{"process": (k if k and k.startswith("IMEB3_") else PROCESS_ALIAS.get((k or "").upper())), "raw": v} for k, v in _pairs(g("inputs"))]
    materials = [{"material": MATERIAL_CODES.get((k or "").upper(), k or ""), "raw": v} for k, v in _pairs(g("materials")) if k]
    defects = []
    for k, v in _pairs(g("defects")):
        name = k or v
        qty = int(v) if k and re.fullmatch(r"\d+", v) else None
        defects.append({"name": name, "qty": qty, "raw": v if k else None})
    lots = [{"raw": p} for p in _split(g("lots"))]
    qty = re.sub(r"[^\d]", "", g("gross_output_qty"))
    minutes = re.sub(r"[^\d.]", "", g("work_minutes"))
    wd = row.get("work_date")
    wd = wd.strftime("%Y-%m-%d") if hasattr(wd, "strftime") else g("work_date")
    shift = (g("shift") or "D").upper().replace("주간", "D").replace("야간", "N").replace("DAY", "D").replace("NIGHT", "N")[:1] or "D"
    return JournalPayload(model=g("model").upper(), process=proc, work_date=wd, shift=shift,
                          gross_output_qty=int(qty) if qty else None, work_minutes=int(float(minutes)) if minutes else None,
                          raw_quantity_text=g("gross_output_qty") or None,
                          inputs=inputs, materials=materials, defects=defects, lots=lots,
                          source={"method": method, "file": g("file") or filename, "page": int(g("page")) if g("page").isdigit() else None})


def parse_rows(rows: Iterable[dict], filename: str, method: str) -> list[dict]:
    out = []
    for i, row in enumerate(rows, start=2):
        if not any(str(v).strip() for v in row.values() if v is not None):
            continue
        try:
            p = row_to_payload(row, filename, method)
            out.append({"row": i, "ok": True, "payload": p.model_dump(), "normalized": p.normalized()})
        except Exception as e:  # noqa: BLE001 - 행 단위 오류를 화면에 보여준다
            out.append({"row": i, "ok": False, "error": str(e), "raw": {k: (str(v) if v is not None else "") for k, v in row.items()}})
    return out


def parse_excel(data: bytes, filename: str) -> list[dict]:
    """템플릿(머리글 한 줄)이면 행 단위로, 회사 작업일보 양식 그대로인 xlsx 면 form_workbook 으로 읽는다."""
    import openpyxl.reader.excel as ex
    ex.ExcelReader.read_custom = lambda self: None   # 일부 xlsx 의 문서 속성 오류 우회
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True, max_row=2))
    if not rows:
        return []
    from .form_workbook import looks_like_form, parse_form_workbook
    if looks_like_form(list(rows[0])):
        return parse_form_workbook(data, filename)
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h).strip().lower() if h is not None else "" for h in rows[0]]
    return parse_rows((dict(zip(header, r)) for r in rows[1:]), filename, "EXCEL")


def parse_csv(data: bytes, filename: str) -> list[dict]:
    text = data.decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    reader.fieldnames = [f.strip().lower() for f in (reader.fieldnames or [])]
    return parse_rows(reader, filename, "CSV")


def template_workbook() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "작업일보"
    ws.append(COLUMNS)
    ws.append(["IV", "IMEB3_WINDING", "2026-09-03", "D", 10122, 300, "IMEB3_STITCHING:9.2", "WIRE:8.28; LEAD_PIN:7.22", "vision1=24; weld=8", "", "20260903-1.pdf", 8])
    ws.append(["IV", "IMEB3_CA_PRESS", "2026-09-03", "D", 10189, 300, "IMEB3_WINDING:2026-09-03", "UPPER_CASE:8.10; LOWER_CASE:7.22", "", "", "20260903-1.pdf", 8])
    ws.append(["IV", "IMEB3_FINAL_INSPECTION", "2026-09-03", "D", 3840, 480, "", "", "Marking=3; 미압입=1", "ZH4=2BOX; ZH3=6BOX", "20260903-1공장.pdf", 10])
    ws2 = wb.create_sheet("설명")
    for line in __doc__.strip().splitlines():
        ws2.append([line])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
